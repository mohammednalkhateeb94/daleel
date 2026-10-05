"""ينقل البنود التي صدر فيها قرار من ورقة «للمراجعة»:

- «معتمد»        ← إلى «المقاطع» أو «القوالب»، وتُنقل ملاحظة المراجع إلى «ملاحظات الاعتماد».
- «معتمد بتعديل» ← لا يُنقل حتى يُنفَّذ التعديل في النص ويُكتب في عمود القرار «معتمد»
                    (يُذكر التعديل في الملاحظة فيُنقل معها). السكربت يطبعها لتنفيذها.
- «مرفوض»        ← إلى «المرفوض» مع السبب.

شروط النقل: اسم المراجع والتاريخ، و«موافق» من المختص الأعلم للمقطع الحساس.

    python scripts/promote.py data/register.xlsx
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from register_lib import (FORM, REJ, REV, SEG, TPL, col, put_row, read_rows,  # noqa: E402
                          seg_formula, style_review, style_segments)


def last_row(ws):
    r = ws.max_row
    while r > 1 and not ws.cell(r, 1).value:
        r -= 1
    return r


def main(path):
    wb = openpyxl.load_workbook(path)
    rev, seg, tpl, rej = wb["للمراجعة"], wb["المقاطع"], wb["القوالب"], wb["المرفوض"]
    moved, pending, blocked, done_rows = [], [], [], []
    for row in read_rows(rev, REV):
        dec = (row.get("القرار") or "").strip()
        if not dec:
            continue
        kind, rid, note = row.get("النوع"), row.get("الرقم"), (row.get("التعديل أو الملاحظة") or "").strip()
        reviewer = (row.get("المراجع والتاريخ") or "").strip()
        if dec == "معتمد بتعديل":
            pending.append(f"{rid}: {note or 'لم يُكتب التعديل'}")
            continue
        if not reviewer:
            blocked.append(f"{rid}: ينقصه اسم المراجع والتاريخ")
            continue
        if dec == "مرفوض":
            r = last_row(rej) + 1
            put_row(rej, r, REJ, {"النوع": kind, "الرقم": rid, "الحاجة أو متى يُستخدم": row.get("الحاجة أو متى يُستخدم"),
                                  "الموضع": row.get("الموضع"), "النص": row.get("النص"), "سبب الرفض": note,
                                  "المراجع والتاريخ": reviewer})
            moved.append(f"{rid} ← المرفوض")
            done_rows.append(row["_row"])
            continue
        if dec != "معتمد":
            blocked.append(f"{rid}: قرار غير معروف «{dec}»")
            continue
        if kind == "مقطع":
            sp = (row.get("رأي المختص الأعلم") or "").strip()
            if row.get("حساس؟") == "نعم" and sp != "موافق":
                blocked.append(f"{rid}: مقطع حساس ينقصه «موافق» من المختص الأعلم")
                continue
            notes = [x for x in [note, "مقطع حساس: وافق المختص الأعلم." if row.get("حساس؟") == "نعم" else ""] if x]
            r = last_row(seg) + 1
            vals = {k: row.get(k) for k in ("الحاجة", "الأولوية", "المصدر", "الموضع", "عنوان المسألة", "السؤال في المصدر",
                                             "عبارات مشابهة", "مضمون السؤال", "يجيب هذا المقطع عن", "صياغات المستخدم",
                                             "المستوى", "حساس؟", "لا يُجتزأ", "أسئلة قريبة لا يجيب عنها")}
            vals.update({"رقم": rid, "الحاجة": row.get("الحاجة أو متى يُستخدم"), "النص المعتمد": row.get("النص"),
                         "آيات وأحاديث (توثيق)": row.get("آيات وأحاديث للتوثيق"),
                         "اعتمده (المراجع والتاريخ)": reviewer, "رأي المختص الأعلم": sp or None,
                         "ملاحظات الاعتماد": " ".join(notes) or "اعتُمد دون ملاحظات."})
            put_row(seg, r, SEG, vals)
            seg.cell(r, col(SEG, "صالح للنشر")).value = seg_formula(r)
            seg.cell(r, col(SEG, "صالح للنشر")).fill = FORM
            moved.append(f"{rid} ← المقاطع")
        else:
            r = last_row(tpl) + 1
            put_row(tpl, r, TPL, {"الرمز": rid, "متى يُستخدم": row.get("الحاجة أو متى يُستخدم"),
                                  "النص الذي يراه المستخدم": row.get("النص"),
                                  "مراجعة شرعية؟": "نعم" if "نعم" in str(row.get("ملاحظات الاستخراج") or "") else "لا",
                                  "اعتمده (المراجع والتاريخ)": reviewer, "ملاحظات الاعتماد": note or "اعتُمد دون ملاحظات."},
                    tall=False)
            moved.append(f"{rid} ← القوالب")
        done_rows.append(row["_row"])
    for r in sorted(done_rows, reverse=True):
        rev.delete_rows(r)
    style_review(rev, last_row(rev))
    style_segments(seg, last_row(seg))
    wb.save(path)
    print(f"نُقل {len(moved)}:", *moved, sep="\n  ")
    if pending:
        print("بانتظار تنفيذ التعديل:", *pending, sep="\n  ")
    if blocked:
        print("لم يُنقل:", *blocked, sep="\n  ")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/register.xlsx")
