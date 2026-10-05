"""بنية سجل «دليل» ومساعدات الكتابة عليه. تستخدمها سكربتات الترحيل والنقل والتصدير."""
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
HDR = PatternFill("solid", fgColor="1F2A44")
HF = Font(name=F, bold=True, color="FFFFFF", size=11)
INP = PatternFill("solid", fgColor="FFF8DC")
FORM = PatternFill("solid", fgColor="EDEDED")
OK = PatternFill("solid", fgColor="E8F5E9")
_thin = Side(style="thin", color="C9C9C9")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
WR = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)

# أعمدة ورقة «المقاطع» (المعتمد فقط)
SEG = [("رقم", 8), ("الحاجة", 7), ("الأولوية", 9), ("المصدر", 7), ("الموضع", 26), ("عنوان المسألة", 28),
       ("السؤال في المصدر", 36), ("عبارات مشابهة", 28), ("مضمون السؤال", 32), ("النص المعتمد", 80),
       ("يجيب هذا المقطع عن", 34), ("صياغات المستخدم", 34), ("المستوى", 7), ("حساس؟", 7),
       ("آيات وأحاديث (توثيق)", 30), ("اعتمده (المراجع والتاريخ)", 18), ("رأي المختص الأعلم", 14),
       ("ملاحظات الاعتماد", 40), ("صالح للنشر", 16), ("لا يُجتزأ", 10), ("أسئلة قريبة لا يجيب عنها", 34),
       ("للتفصيل فقط", 34), ("تنبيه المراجع", 34)]

# أعمدة ورقة «للمراجعة» (مقاطع وقوالب تنتظر قراراً)
REV = [("النوع", 7), ("الرقم", 8), ("الحاجة أو متى يُستخدم", 16), ("الأولوية", 9), ("المصدر", 7), ("الموضع", 26),
       ("عنوان المسألة", 24), ("السؤال في المصدر", 30), ("عبارات مشابهة", 24), ("مضمون السؤال", 26),
       ("النص", 80), ("يجيب هذا المقطع عن", 32), ("صياغات المستخدم", 30), ("المستوى", 7), ("حساس؟", 7),
       ("آيات وأحاديث للتوثيق", 28), ("ملاحظات الاستخراج", 34), ("ما المطلوب من المختص", 40),
       ("القرار", 13), ("التعديل أو الملاحظة", 36), ("المراجع والتاريخ", 18), ("رأي المختص الأعلم", 14),
       ("لا يُجتزأ", 10), ("أسئلة قريبة لا يجيب عنها", 34), ("للتفصيل فقط", 34), ("تنبيه المراجع", 34)]
REV_INPUT = {"القرار", "التعديل أو الملاحظة", "المراجع والتاريخ", "رأي المختص الأعلم"}

TPL = [("الرمز", 8), ("متى يُستخدم", 28), ("النص الذي يراه المستخدم", 80), ("مراجعة شرعية؟", 12),
       ("اعتمده (المراجع والتاريخ)", 18), ("ملاحظات الاعتماد", 36)]

REJ = [("النوع", 7), ("الرقم", 8), ("الحاجة أو متى يُستخدم", 16), ("الموضع", 26), ("النص", 70),
       ("سبب الرفض", 40), ("المراجع والتاريخ", 18)]

GAP = [("رقم", 8), ("الحاجة", 7), ("الأولوية", 9), ("ما المطلوب", 44), ("المصدر المرشح", 44), ("ملاحظات", 30)]


def col(cols, name):
    return [c for c, _ in cols].index(name) + 1


def header(ws, cols):
    ws.sheet_view.rightToLeft = True
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.fill, c.font, c.border = HDR, HF, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center", readingOrder=2)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"


def put_row(ws, r, cols, values: dict, inputs=(), tall=True):
    for i, (name, _) in enumerate(cols, 1):
        c = ws.cell(r, i, values.get(name))
        c.font, c.alignment, c.border = Font(name=F, size=10, bold=(i <= 2)), WR, BOX
        if name in inputs:
            c.fill = INP
    ws.row_dimensions[r].height = 150 if tall and values.get("النص") or values.get("النص المعتمد") else 34


def read_rows(ws, cols):
    names = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    out = []
    for r in range(2, ws.max_row + 1):
        row = {n: ws.cell(r, i + 1).value for i, n in enumerate(names) if n}
        if any(v not in (None, "") for v in row.values()):
            row["_row"] = r
            out.append(row)
    return out


def dv(ws, rng, items):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)


def seg_formula(r):
    c = {n: get_column_letter(col(SEG, n)) for n, _ in SEG}
    src_ok = f'IFERROR(INDEX(المصادر!$F$2:$F$30,MATCH({c["المصدر"]}{r},المصادر!$A$2:$A$30,0)),"")'
    return (f'=IF({c["رقم"]}{r}="","",IF(OR({c["المستوى"]}{r}="ج",{c["المستوى"]}{r}="د"),"المستوى لا يُنشر",'
            f'IF({src_ok}<>"نعم","المصدر غير معتمد",'
            f'IF(OR({c["الحاجة"]}{r}="",{c["الموضع"]}{r}="",{c["النص المعتمد"]}{r}="",{c["يجيب هذا المقطع عن"]}{r}="",{c["اعتمده (المراجع والتاريخ)"]}{r}=""),"بيانات ناقصة",'
            f'IF(AND({c["حساس؟"]}{r}="نعم",{c["رأي المختص الأعلم"]}{r}<>"موافق"),"ينقصه رأي المختص","صالح للنشر")))))')


def style_segments(ws, last):
    letter = get_column_letter(col(SEG, "صالح للنشر"))
    ws.conditional_formatting.add(f"{letter}2:{letter}200", FormulaRule(formula=[f'${letter}2="صالح للنشر"'], fill=OK, font=Font(name=F, color="006100", bold=True)))
    n = get_column_letter(col(SEG, "حساس؟"))
    ws.conditional_formatting.add(f"{n}2:{n}200", FormulaRule(formula=[f'${n}2="نعم"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SEG))}{max(last, 2)}"


def style_review(ws, last):
    d = get_column_letter(col(REV, "القرار"))
    dv(ws, f"{d}2:{d}300", ["معتمد", "معتمد بتعديل", "مرفوض"])
    s = get_column_letter(col(REV, "رأي المختص الأعلم"))
    dv(ws, f"{s}2:{s}300", ["موافق", "موافق بتعديل", "غير موافق"])
    n = get_column_letter(col(REV, "حساس؟"))
    ws.conditional_formatting.add(f"{n}2:{n}300", FormulaRule(formula=[f'${n}2="نعم"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REV))}{max(last, 2)}"
